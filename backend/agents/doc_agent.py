# -*- coding: utf-8 -*-
"""
文档信息Agent（Kimi K3）—— 系统知识库，超长文档处理专项Agent

架构文档来源：第4章 4.3
  定位：系统知识库，超长文档处理专项Agent
  核心能力：PDF/Word/MD/代码库超长解析、信息抽取、结构化摘要、资料对比
  约束：禁止输出原始长文本，仅输出摘要+引用片段，降低全局Token消耗
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from backend.agents.base import AGENT_DOC, CAPABILITY_DOC_PARSE, BaseAgent
from backend.bus.message import Message
from backend.bus.state_machine import TaskState
from backend.services.model_client import ModelUnavailable
from backend.utils.constants import AGENT_DOC, AGENT_TEMPERATURES

DOC_SYSTEM_PROMPT = """你是「文档信息Agent」（Kimi K3），本地化多Agent协同开发工作台的知识库与超长文档处理专项角色。

【硬约束 · 不可违反】
1. 严禁输出原始长文本。任何情况下只输出「结构化摘要 + 少量引用片段」。
2. 每个引用片段不得超过 200 字，单次回复中引用片段总量不得超过原文的 10%。
3. 你只能读取当前会话目录内的文档，不得读取目录外文件，不得执行任何文件写入或系统命令。
4. 无法确定的内容必须标注「未在文档中找到」，禁止编造。

【输出格式 · 必须是合法 JSON】
{
  "doc_name": "文件相对路径",
  "doc_type": "pdf/docx/md/txt/csv/code",
  "summary": "结构化摘要（Markdown，分点，300-800字）",
  "outline": ["一级小节1", "一级小节2"],
  "key_points": ["要点1", "要点2"],
  "snippets": [ {"locator": "第3章/第2段", "quote": "不超过200字的引用片段"} ],
  "comparison": "当要求多文档对比时给出对比结论，否则留空",
  "token_saving_note": "说明本次如何通过摘要替代原文降低Token消耗"
}
"""

# 单文件读取上限：超长文档只抽取前 N 字符进入上下文（第4章 4.3 降低全局Token消耗）
MAX_DOC_CHARS = 120_000
SNIPPET_LIMIT = 200


class DocAgent(BaseAgent):
    agent_role = AGENT_DOC
    model_name = "Kimi K3"

    def __init__(self, ctx):
        super().__init__(ctx)
        self.system_prompt = DOC_SYSTEM_PROMPT

    # ==================================================================
    # 文档读取（仅当前会话目录内）
    # ==================================================================
    def read_document(self, relative_path: str) -> tuple[bool, str, str]:
        """返回 (ok, doc_type, text)。支持 pdf / docx / md / txt / csv / 代码。"""
        self.require_capability(CAPABILITY_DOC_PARSE)
        guard = self.ctx.file_guard
        target = guard.resolve(relative_path)
        if not target.exists() or target.is_dir():
            return False, "", f"文档不存在：{relative_path}"

        ext = target.suffix.lower()
        try:
            if ext == ".pdf":
                return True, "pdf", self._read_pdf(target)
            if ext == ".docx":
                return True, "docx", self._read_docx(target)
            if ext == ".csv":
                return True, "csv", self._read_csv(target)
            if ext in (".xlsx", ".xls"):
                return True, "excel", self._read_excel(target)
            text = target.read_text(encoding="utf-8", errors="replace")
            return True, "md" if ext in (".md", ".markdown") else "txt", text
        except Exception as exc:  # noqa: BLE001
            self.ctx.logger.exception_log(
                error_code="DOC_READ_FAILED", message=f"读取文档失败 {relative_path}: {exc}",
                agent_role=self.agent_role,
            )
            return False, "", f"文档解析失败：{exc}"

    @staticmethod
    def _read_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError:  # pragma: no cover
            return "（未安装 pypdf，无法解析 PDF。请在后端安装 pypdf）"
        reader = PdfReader(str(path))
        out: list[str] = []
        for page in reader.pages[:300]:
            try:
                out.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001
                continue
        return "\n".join(out)

    @staticmethod
    def _read_docx(path: Path) -> str:
        try:
            import docx  # python-docx
        except ImportError:  # pragma: no cover
            return "（未安装 python-docx，无法解析 Word 文档）"
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables[:50]:
            for row in table.rows[:200]:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)

    @staticmethod
    def _read_csv(path: Path) -> str:
        raw = path.read_text(encoding="utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(raw)))
        head = rows[:80]
        return "\n".join(" | ".join(r) for r in head)

    @staticmethod
    def _read_excel(path: Path) -> str:
        try:
            import openpyxl  # type: ignore
        except ImportError:
            return "（未安装 openpyxl，无法解析 Excel。建议另存为 CSV 后再解析）"
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in list(wb.worksheets)[:5]:
            lines.append(f"[sheet] {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 80:
                    break
                lines.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(lines)

    # ==================================================================
    # 摘要抽取主流程
    # ==================================================================
    async def analyze(self, task: TaskState, relative_path: str, *, focus: str = "",
                      compare_with: list[str] | None = None) -> dict:
        self.require_capability(CAPABILITY_DOC_PARSE)
        task.bump_iteration()
        self.ctx.runtime.router.check_timeout(task)

        ok, doc_type, text = self.read_document(relative_path)
        if not ok:
            self.think(task, "think", f"文档无法读取：{text}")
            return {"ok": False, "error": text, "doc_name": relative_path}

        self.think(task, "read", f"解析文档 {relative_path}（{doc_type}，{len(text)} 字符）")

        # 第4章 4.3 约束：不把原始长文本交给后续环节，只做受限抽取
        clipped = text[:MAX_DOC_CHARS]
        clipped_note = ""
        if len(text) > MAX_DOC_CHARS:
            clipped_note = f"\n（原文共 {len(text)} 字符，本次仅抽取前 {MAX_DOC_CHARS} 字符用于摘要，已按约束丢弃其余原文）"

        compare_blocks = []
        for other in (compare_with or [])[:3]:
            ok2, t2, txt2 = self.read_document(other)
            if ok2:
                compare_blocks.append(f"【对比文档：{other}】\n{txt2[:30000]}")

        user_content = "\n\n".join(x for x in [
            f"【待解析文档】{relative_path}（类型 {doc_type}）{clipped_note}",
            f"【解析聚焦点】{focus}" if focus else "",
            "【文档内容抽取】\n" + clipped,
            "\n\n".join(compare_blocks),
        ] if x)

        messages = [
            {"role": "system", "content": self.system_prompt},
            {"role": "user", "content": user_content},
        ]

        try:
            response = await self.call_model(
                # 【需求点 Bug4】文档/信息检索子 Agent 温度 0.25（区间 0.2~0.3）
                task, messages, temperature=AGENT_TEMPERATURES[AGENT_DOC],
                max_tokens=4000, expect_json=True,
                step_label="Kimi K3 解析长文档并生成结构化摘要",
            )
        except ModelUnavailable as exc:
            return {"ok": False, "error": f"文档信息能力当前不可用：{exc}", "doc_name": relative_path}

        try:
            data = self.parse_json(response.text)
        except ValueError:
            data = {"summary": response.text[:3000], "snippets": [], "key_points": []}

        if not isinstance(data, dict):
            data = {"summary": str(data), "snippets": [], "key_points": []}

        # 强制收敛引用片段长度（第4章 4.3 硬约束，后端兜底裁剪）
        snippets = []
        for s in list(data.get("snippets") or [])[:20]:
            if isinstance(s, dict):
                quote = str(s.get("quote") or "")[:SNIPPET_LIMIT]
                snippets.append({"locator": str(s.get("locator") or "未知位置"), "quote": quote})
            elif isinstance(s, str):
                snippets.append({"locator": "未知位置", "quote": s[:SNIPPET_LIMIT]})
        data["snippets"] = snippets
        data["doc_name"] = relative_path
        data["doc_type"] = doc_type
        data["ok"] = True
        data.setdefault("token_saving_note",
                        f"仅输出摘要+{len(snippets)} 条片段，原文 {len(text)} 字符未进入后续上下文")
        self.think(task, "write", f"输出结构化摘要（引用片段 {len(snippets)} 条，均已裁剪至 {SNIPPET_LIMIT} 字内）")
        return data

    # ==================================================================
    def build_output(self, data: dict) -> str:
        """转成给调度Agent的摘要文本（结构化，不含原文）。"""
        if not data.get("ok"):
            return f"文档解析失败：{data.get('error', '未知原因')}"
        lines = [
            f"## 文档摘要：{data.get('doc_name')}",
            "",
            str(data.get("summary") or "").strip(),
        ]
        points = data.get("key_points") or []
        if points:
            lines += ["", "## 关键要点"]
            lines += [f"- {p}" for p in points[:20]]
        outline = data.get("outline") or []
        if outline:
            lines += ["", "## 结构大纲"]
            lines += [f"- {o}" for o in outline[:30]]
        if data.get("comparison"):
            lines += ["", "## 对比结论", str(data["comparison"])]
        if data.get("snippets"):
            lines += ["", "## 引用片段"]
            lines += [f"- [{s['locator']}] {s['quote']}" for s in data["snippets"][:10]]
        if data.get("token_saving_note"):
            lines += ["", f"> Token 说明：{data['token_saving_note']}"]
        return "\n".join(lines)

    # ==================================================================
    async def handle(self, msg: Message, task: TaskState) -> Message | None:
        self.require_capability(CAPABILITY_DOC_PARSE)
        meta = msg.metadata
        path = str(meta.get("doc_path") or meta.get("path") or "").strip()
        focus = str(meta.get("instruction") or msg.content_text())
        if not path:
            return self.error_message(task, "DOC_PATH_MISSING",
                                      "文档信息Agent 未收到文档路径（metadata.doc_path）")
        data = await self.analyze(task, path, focus=focus,
                                  compare_with=list(meta.get("compare_with") or []))
        if not data.get("ok"):
            return self.error_message(task, "DOC_ANALYZE_FAILED", str(data.get("error")))
        return self.result_message(task, self.build_output(data), metadata={
            "doc_name": data.get("doc_name"),
            "doc_type": data.get("doc_type"),
            "snippets": data.get("snippets"),
            "key_points": data.get("key_points"),
            "no_raw_text": True,
        })


__all__ = ["DocAgent", "DOC_SYSTEM_PROMPT", "MAX_DOC_CHARS"]
